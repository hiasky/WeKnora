#!/usr/bin/env python3
"""Create a restrained, readable PPTX from a JSON document on stdin."""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


MAX_SLIDES = 50
SAFE_NAME = re.compile(r"^[^/\\\x00-\x1f]+\.pptx$", re.IGNORECASE)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)
ACCENT = RGBColor(7, 192, 95)
WHITE = RGBColor(255, 255, 255)


def fail(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    raise SystemExit(2)


def text(value: object, limit: int) -> str:
    return str(value or "").strip()[:limit]


def add_title(slide, title: str, subtitle: str) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = INK
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.5))
    frame = title_box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(38)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    paragraph.alignment = PP_ALIGN.LEFT
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.94), Inches(3.75), Inches(10.8), Inches(0.8))
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(209, 213, 219)
    bar = slide.shapes.add_shape(1, Inches(0.9), Inches(1.78), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_content(slide, title: str, bullets: list[str], number: int) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.4), Inches(0.75))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = INK

    rule = slide.shapes.add_shape(1, Inches(0.76), Inches(1.42), Inches(0.75), Inches(0.055))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()

    body = slide.shapes.add_textbox(Inches(0.85), Inches(1.75), Inches(11.25), Inches(4.8))
    frame = body.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(22 if len(bullets) <= 5 else 19)
        paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(14)
        paragraph.line_spacing = 1.12

    footer = slide.shapes.add_textbox(Inches(11.6), Inches(7.0), Inches(0.75), Inches(0.25))
    fp = footer.text_frame.paragraphs[0]
    fp.text = str(number)
    fp.font.size = Pt(10)
    fp.font.color.rgb = MUTED
    fp.alignment = PP_ALIGN.RIGHT


def main() -> None:
    try:
        payload = json.load(sys.stdin)
    except Exception as exc:
        fail(f"stdin must contain valid JSON: {exc}")

    title = text(payload.get("title"), 160)
    if not title:
        fail("title is required")
    raw_slides = payload.get("slides")
    if not isinstance(raw_slides, list) or not 1 <= len(raw_slides) <= MAX_SLIDES:
        fail(f"slides must contain 1 to {MAX_SLIDES} items")
    filename = text(payload.get("output_file") or "presentation.pptx", 180)
    if not SAFE_NAME.fullmatch(filename) or filename in {".", ".."}:
        fail("output_file must be a plain .pptx filename")

    output_root = Path(os.environ.get("WEKNORA_SKILL_OUTPUT_DIR", "/workspace/output"))
    output_root.mkdir(parents=True, exist_ok=True)
    destination = output_root / filename

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    subtitle = text(payload.get("subtitle"), 240)
    author = text(payload.get("author"), 120)
    add_title(cover, title, " · ".join(part for part in (subtitle, author) if part))

    for number, raw in enumerate(raw_slides, start=1):
        if not isinstance(raw, dict):
            fail(f"slide {number} must be an object")
        slide_title = text(raw.get("title"), 140) or f"Slide {number}"
        raw_bullets = raw.get("bullets") or []
        if not isinstance(raw_bullets, list):
            fail(f"slide {number} bullets must be an array")
        bullets = [text(item, 420) for item in raw_bullets if text(item, 420)] or [" "]
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        add_content(slide, slide_title, bullets[:10], number)
        notes = text(raw.get("notes"), 4000)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    deck.save(destination)
    print(json.dumps({
        "ok": True,
        "artifact_type": "presentation",
        "path": str(destination),
        "file_name": destination.name,
        "slides": len(deck.slides),
        "bytes": destination.stat().st_size,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
